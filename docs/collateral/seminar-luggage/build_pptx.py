"""BondEx 講演スライド PPTX ビルド（build.mjs の計測結果 manifest.json を読む）

  python3 build_pptx.py            写真あり版（build/photo）
  python3 build_pptx.py --type     タイポグラフィ版（build/type）

テキストはすべて編集可能なテキストボックス、写真は「図の変更」で差し替え可能な画像として配置する。
画像に文字は焼き込まない。
"""
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

DIR = Path(__file__).resolve().parent
TYPE_ONLY = "--type" in sys.argv
BUILD = DIR / "build" / ("type" if TYPE_ONLY else "photo")
SPEC = json.loads((DIR / "slides.json").read_text())
MAN = json.loads((BUILD / "manifest.json").read_text())
NOTES_PATH = DIR / "notes.json"
NOTES = {n["n"]: n for n in json.loads(NOTES_PATH.read_text())["slides"]} if NOTES_PATH.exists() else {}

W, H = SPEC["canvas"]["w"], SPEC["canvas"]["h"]
EMU_PX = 12192000 // W  # 16:9 / 13.333in を 1920px として扱う
COL = {k: RGBColor.from_string(v.lstrip("#")) for k, v in SPEC["colors"].items()}
FONTS = SPEC["fonts"]
TYPE_SCALE = 1.14
# Chromium と PowerPoint で「行ボックス内のベースライン位置」が違うぶんの補正（em 単位、LibreOffice 描画で校正）
BASELINE_DY = {"ja": -0.075, "jaRegular": -0.075, "en": 0.04}
# BASELINE_DY を校正した時の行送り（em）と、フォントの hhea ascent / descent（em）
LH_REF = {"ja": 1.42, "jaRegular": 1.42, "en": 1.04}
METRICS = {"ja": (1.16, 0.288), "jaRegular": (1.16, 0.288), "en": (0.96875, 0.2412)}


def px(v):
    return Emu(int(round(v * EMU_PX)))


def pt(v):  # 1px = 0.5pt
    return Pt(v * 0.5)


def set_font(run, key, size_px, color, track_em=0.0):
    f = FONTS[key]
    run.font.size = pt(size_px)
    run.font.color.rgb = color
    run.font.bold = False
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", f["pptx"])
    if track_em:
        rpr.set("spc", str(int(round(track_em * size_px * 0.5 * 100))))
    rpr.set("lang", "ja-JP")
    rpr.set("altLang", "en-US")


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if name:
        tb.name = name
    return tb, tf


def exact_spacing(p, line_px):
    p.line_spacing = pt(line_px)
    p.space_before = p.space_after = 0


def add_picture(slide, ph, x, y, w, h, name):
    pic = slide.shapes.add_picture(ph["path"], px(x), px(y), px(w), px(h))
    c = ph["crop"]
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = c["l"], c["t"], c["r"], c["b"]
    pic.name = name
    if ph.get("opacity", 1) < 1:
        blip = pic._element.find(".//" + qn("a:blip"))
        amf = etree.SubElement(blip, qn("a:alphaModFix"))
        amf.set("amt", str(int(ph["opacity"] * 100000)))
    return pic


def add_scrim(slide, scrim):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(W), px(H))
    shp.name = "Scrim（写真の上の黒グラデーション）"
    shp.line.fill.background()
    sppr = shp._element.spPr
    for el in sppr.findall(qn("a:solidFill")) + sppr.findall(qn("a:noFill")):
        sppr.remove(el)
    grad = etree.Element(qn("a:gradFill"), rotWithShape="1")
    gsl = etree.SubElement(grad, qn("a:gsLst"))
    for pos, alpha in scrim["stops"]:
        gs = etree.SubElement(gsl, qn("a:gs"), pos=str(int(pos * 1000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    # OOXML の lin ang は「時計回り、0=左→右」。CSS の向きに合わせて変換済みの値を使う
    etree.SubElement(grad, qn("a:lin"), ang=str(int(scrim["ang"] * 60000)), scaled="0")
    geom = sppr.find(qn("a:prstGeom"))
    geom.addnext(grad)
    # 影・枠などテーマ由来の装飾を消す
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)
    ln = sppr.find(qn("a:ln"))
    if ln is None:
        ln = etree.SubElement(sppr, qn("a:ln"))
    for ch in list(ln):
        ln.remove(ch)
    etree.SubElement(ln, qn("a:noFill"))
    return shp


def headline(slide, s, m):
    t = s["text"]
    key = t["font"]
    size = t["size"] * (TYPE_SCALE if TYPE_ONLY and key != "en" and s["layout"] not in ("diagram", "cover") else 1)
    line = size * t["lh"]
    box = m["measured"]["txt"]
    hang = size * 0.5
    any_hang = any(runs[0]["t"].startswith("「") for runs in t["lines"])
    x0 = box["x"] - (hang if any_hang else 0)
    dy = BASELINE_DY[key] * size
    tb, tf = textbox(slide, x0, box["y"] + dy, box["w"] + hang + 80, box["h"], name="見出し")
    for i, runs in enumerate(t["lines"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        exact_spacing(p, line)
        p.alignment = PP_ALIGN.LEFT
        if any_hang and not runs[0]["t"].startswith("「"):
            p._p.get_or_add_pPr().set("marL", str(int(px(hang))))
        for r in runs:
            run = p.add_run()
            run.text = r["t"]
            set_font(run, key, size, COL[r.get("c", "fg")], t["track"])


def number_slide(slide, s, m):
    t = s["text"]
    meas = m["measured"]
    for k, text in (("pre", t["pre"]), ("post", t["post"])):
        b = meas[k]
        tb, tf = textbox(slide, 0, b["y"] + BASELINE_DY["ja"] * t["small"], W, b["h"], name="数字の前後")
        p = tf.paragraphs[0]
        exact_spacing(p, t["small"] * 1.3)
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        set_font(run, "ja", t["small"], RGBColor(0xC7, 0xC7, 0xC7), -0.01)  # 白 78% 相当
    # 数字行：約 / 50 / 万円 を別々のボックスにし、ブラウザで計測したベースラインに合わせる
    for r, b in zip(t["num"], meas["numRuns"]):
        key, size = r["font"], r["size"]
        lh = LH_REF[key]
        a, d = METRICS[key]
        top = b["baseline"] - ((lh - (a + d)) / 2 + a) * size + BASELINE_DY[key] * size
        tb, tf = textbox(slide, b["x"], top, b["w"] + 60, size * lh, name="数字")
        p = tf.paragraphs[0]
        exact_spacing(p, size * lh)
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = r["t"]
        set_font(run, key, size, COL["accent"], -0.04 if key == "en" else -0.02)
    b = meas["note"]
    tb, tf = textbox(slide, b["x"], b["y"], b["w"] + 60, b["h"], name="注記")
    p = tf.paragraphs[0]
    exact_spacing(p, b["h"])
    run = p.add_run()
    run.text = t["note"]
    set_font(run, "jaRegular", 22, COL["muted"], 0.01)


def qpts(q, n=64):
    (x0, y0), (cx, cy), (x1, y1) = q
    return [((1 - t) ** 2 * x0 + 2 * (1 - t) * t * cx + t * t * x1, (1 - t) ** 2 * y0 + 2 * (1 - t) * t * cy + t * t * y1) for t in (i / n for i in range(n + 1))]


def line_style(shp, color_hex, width_px, alpha=1.0, dash=None):
    """線の色・太さ・透明度・点線（丸端）を DrawingML で直接指定する"""
    sppr = shp._element.spPr
    for tag in ("a:ln", "a:solidFill", "a:noFill", "a:gradFill"):
        for el in sppr.findall(qn(tag)):
            sppr.remove(el)
    etree.SubElement(sppr, qn("a:noFill"))
    ln = etree.SubElement(sppr, qn("a:ln"), w=str(int(px(width_px))), cap="rnd")
    fill = etree.SubElement(ln, qn("a:solidFill"))
    clr = etree.SubElement(fill, qn("a:srgbClr"), val=color_hex.lstrip("#"))
    if alpha < 1:
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
    if dash:
        cd = etree.SubElement(ln, qn("a:custDash"))
        etree.SubElement(cd, qn("a:ds"), d=str(int(dash[0] / width_px * 100000)), sp=str(int(dash[1] / width_px * 100000)))
    etree.SubElement(ln, qn("a:round"))
    style = shp._element.find(qn("p:style"))
    if style is not None:
        shp._element.remove(style)


def diagram(slide, s):
    g = s["diagram"]
    hexc = SPEC["colors"]
    for i, p in enumerate(g["paths"]):
        pts = qpts(p["q"]) if "q" in p else [tuple(v) for v in p["pts"]]
        fb = slide.shapes.build_freeform(px(pts[0][0]), px(pts[0][1]), scale=1.0)
        fb.add_line_segments([(px(x), px(y)) for x, y in pts[1:]], close=False)
        shp = fb.convert_to_shape()
        shp.name = f"図・線{i + 1}"
        line_style(shp, hexc[p["c"]], p["w"], p.get("alpha", 1), p.get("dash"))
    for o in g["dots"]:
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(o["x"] - o["r"]), px(o["y"] - o["r"]), px(o["r"] * 2), px(o["r"] * 2))
        d.name = "図・点"
        d.fill.solid()
        d.fill.fore_color.rgb = COL[o["c"]]
        d.line.fill.background()
        style = d._element.find(qn("p:style"))
        if style is not None:
            d._element.remove(style)
    for l in g["labels"]:
        size, key = l["size"], l["font"]
        x = l["x"] - 900 if l["align"] == "center" else l["x"] - 1800 if l["align"] == "right" else l["x"]
        tb, tf = textbox(slide, x, l["y"] + BASELINE_DY[key] * size, 1800, size * 1.42, name="図・ラベル")
        p = tf.paragraphs[0]
        exact_spacing(p, size * 1.42)
        p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(l["align"], PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = l["t"]
        set_font(run, key, size, COL[l["c"]], -0.01)


def cover(slide, s, m):
    meas, c = m["measured"], s["cover"]
    b = meas["clogo"]
    pic = slide.shapes.add_picture(str(DIR / "assets/bondex-logo-red.png"), px(b["x"]), px(b["y"]), px(b["w"]), px(b["h"]))
    pic.name = "BondEx ロゴ"
    headline(slide, s, m)
    b = meas["tag"]
    tb, tf = textbox(slide, b["x"], b["y"] + BASELINE_DY["en"] * 28, b["w"] + 80, b["h"], name="タグライン")
    p = tf.paragraphs[0]
    exact_spacing(p, 28 * 1.42)
    run = p.add_run()
    run.text = c["tagline"]
    set_font(run, "en", 28, COL["muted"], 0.2)
    b = meas["qr"]
    pic = slide.shapes.add_picture(str(DIR / c["qr"]), px(b["x"]), px(b["y"]), px(b["w"]), px(b["h"]))
    pic.name = f"QRコード（{c['qrUrl']}）"
    pic.click_action.hyperlink.address = c["qrUrl"]
    for key, font_key, size, color, track in (("spkorg", "jaRegular", 26, COL["muted"], 0.02), ("spkname", "ja", 44, COL["fg"], 0.06)):
        b = meas[key]
        tb, tf = textbox(slide, b["x"], b["y"] + BASELINE_DY[font_key] * size, b["w"] + 120, b["h"], name="登壇者")
        p = tf.paragraphs[0]
        exact_spacing(p, size * 1.42)
        run = p.add_run()
        run.text = c["speaker"]["org" if key == "spkorg" else "name"]
        set_font(run, font_key, size, color, track)
    b = meas["qrlab"]
    tb, tf = textbox(slide, b["x"] + b["w"] - 600, b["y"] + BASELINE_DY["jaRegular"] * 24, 600, b["h"], name="QRの説明")
    p = tf.paragraphs[0]
    exact_spacing(p, 24 * 1.42)
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = c["qrLabel"]
    set_font(run, "jaRegular", 24, COL["muted"])


def notes_text(n):
    d = NOTES.get(n)
    if not d:
        return ""
    return "\n\n".join([f"［{d['time']}］{d['purpose']}", d["script"].replace("\n\n", "\n"), f"◆ 強調：{d['emphasis']}", f"→ つなぎ：{d['transition']}"])


prs = Presentation()
prs.slide_width, prs.slide_height = px(W), px(H)
blank = prs.slide_layouts[6]
split = MAN["split"]

for s, m in zip(SPEC["slides"], MAN["slides"]):
    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = COL["bg"]
    if s["layout"] == "full":
        if m.get("photo"):
            add_picture(slide, m["photo"], 0, 0, W, H, "写真（右クリック→図の変更で差し替え）")
            if s.get("scrim"):
                add_scrim(slide, MAN["scrims"][s["scrim"]])
        headline(slide, s, m)
    elif s["layout"] == "split":
        hw = (W - split["gap"]) / 2
        if m.get("photo"):
            add_picture(slide, m["photo"], 0, 0, hw, split["h"], "写真・左（図の変更で差し替え）")
        if m.get("photo2"):
            add_picture(slide, m["photo2"], hw + split["gap"], 0, hw, split["h"], "写真・右（図の変更で差し替え）")
        headline(slide, s, m)
    elif s["layout"] == "cover":
        cover(slide, s, m)
    elif s["layout"] == "diagram":
        diagram(slide, s)
        headline(slide, s, m)
    elif s["layout"] == "number":
        if m.get("photo"):
            add_picture(slide, m["photo"], 0, 0, W, H, "背景写真（任意）")
        number_slide(slide, s, m)
    if s.get("logo") and m["measured"].get("logo"):
        b = m["measured"]["logo"]
        pic = slide.shapes.add_picture(str(DIR / "assets/bondex-logo-red.png"), px(b["x"]), px(b["y"]), px(b["w"]), px(b["h"]))
        pic.name = "BondEx ロゴ"
        blip = pic._element.find(".//" + qn("a:blip"))
        etree.SubElement(blip, qn("a:alphaModFix"), amt="92000")
    txt = notes_text(s["n"])
    if txt:
        slide.notes_slide.notes_text_frame.text = txt

name = "BondEx_Seminar_TypeOnly.pptx" if TYPE_ONLY else MAN["pdfName"].replace(".pdf", ".pptx")
out = DIR / "out" / name
prs.save(out)
print("PPTX:", out.relative_to(DIR))
